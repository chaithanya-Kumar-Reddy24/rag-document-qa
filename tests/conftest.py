"""Shared pytest fixtures.

No test calls a real API: OpenAI embeddings and chat models are replaced with
small deterministic fakes so the whole suite runs offline and for free.
"""

from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any

import pytest
from docx import Document as DocxDocument
from langchain_core.documents import Document
from langchain_core.embeddings import Embeddings
from langchain_core.language_models.fake_chat_models import FakeListChatModel
from reportlab.pdfgen import canvas

from src.config import Settings
from src.rag_chain import AnswerGenerator, QueryRewriter
from src.retriever import DocumentRetriever
from src.vector_store import VectorStoreManager


class FakeEmbeddings(Embeddings):
    """Deterministic, dependency-free stand-in for OpenAIEmbeddings.

    Turns text into a small vector using word hashes, so semantically similar
    text (shared words) lands closer together under cosine similarity - good
    enough to exercise retrieval logic without any network calls.
    """

    def __init__(self, dimensions: int = 32) -> None:
        self.dimensions = dimensions

    def _vector(self, text: str) -> list[float]:
        vector = [0.0] * self.dimensions
        for word in text.lower().split():
            index = int(hashlib.sha256(word.encode()).hexdigest(), 16) % self.dimensions
            vector[index] += 1.0
        norm = sum(component**2 for component in vector) ** 0.5
        return [component / norm for component in vector] if norm else vector

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        return [self._vector(text) for text in texts]

    def embed_query(self, text: str) -> list[float]:
        return self._vector(text)


@pytest.fixture
def tmp_settings(tmp_path: Path) -> Settings:
    return Settings(
        openai_api_key="sk-test-key",
        chroma_persist_directory=tmp_path / "chroma_db",
        upload_directory=tmp_path / "uploads",
        chunk_size=200,
        chunk_overlap=20,
        top_k=4,
        relevance_threshold=0.05,
    )


@pytest.fixture
def fake_embeddings() -> FakeEmbeddings:
    return FakeEmbeddings()


@pytest.fixture
def vector_store(tmp_settings: Settings, fake_embeddings: FakeEmbeddings) -> VectorStoreManager:
    return VectorStoreManager(tmp_settings, fake_embeddings)


@pytest.fixture
def retriever(vector_store: VectorStoreManager, tmp_settings: Settings) -> DocumentRetriever:
    return DocumentRetriever(vector_store, tmp_settings.top_k, tmp_settings.relevance_threshold)


def make_fake_llm(responses: list[str]) -> FakeListChatModel:
    return FakeListChatModel(responses=responses)


@pytest.fixture
def make_generator():
    def _make(responses: list[str]) -> AnswerGenerator:
        return AnswerGenerator(make_fake_llm(responses))

    return _make


@pytest.fixture
def make_reranker():
    def _make(responses: list[str]):
        from src.reranker import LLMReranker

        return LLMReranker(make_fake_llm(responses))

    return _make


@pytest.fixture
def make_rewriter():
    def _make(responses: list[str]) -> QueryRewriter:
        return QueryRewriter(make_fake_llm(responses))

    return _make


@pytest.fixture
def sample_docx(tmp_path: Path) -> Path:
    path = tmp_path / "sample.docx"
    document = DocxDocument()
    document.add_paragraph("Supervised learning uses labeled training data.")
    document.add_paragraph("Unsupervised learning finds patterns without labels.")
    document.save(str(path))
    return path


@pytest.fixture
def sample_pdf(tmp_path: Path) -> Path:
    path = tmp_path / "sample.pdf"
    pdf = canvas.Canvas(str(path))
    pdf.drawString(72, 720, "Python is a popular programming language for data science.")
    pdf.showPage()
    pdf.drawString(72, 720, "Machine learning models are trained on data.")
    pdf.showPage()
    pdf.save()
    return path


def make_chunk(text: str, **metadata: Any) -> Document:
    base = {"source": "doc.txt", "file_type": "txt", "doc_id": "abc123", "chunk_id": "doc_01"}
    base.update(metadata)
    return Document(page_content=text, metadata=base)
